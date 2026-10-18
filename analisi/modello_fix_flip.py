"""
Modello fix & flip di un lotto d'asta, con formule vive (skill moolaski-financial-analyst).

Uso: python3 analisi/modello_fix_flip.py <lotto>   (lotti in analisi/lotti_fix_flip.py)
Ogni valore porta la fonte: dato (avviso, perizia, OMI) o IPOTESI dichiarata.
Verifica: python3 .claude/skills/moolaski-financial-analyst/scripts/check_workbook.py <file.xlsx>
"""
import datetime as dt
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as L

BLU, NERO, VERDE = Font(color="0000FF"), Font(color="000000"), Font(color="008000")
BOLD = Font(bold=True); TIT = Font(bold=True, size=14)
TOP = Border(top=Side(style="thin"))
EUR = '#,##0" €";[Red]-#,##0" €"'; PCT = '0.0%'; MQ = '#,##0.0" mq"'; NUM = '#,##0'
def costruisci(lotto, percorso):
 wb = Workbook()
 cover = wb.active; cover.title = "Cover"
 inp = wb.create_sheet("Inputs"); cf = wb.create_sheet("CF"); out = wb.create_sheet("Outputs"); chk = wb.create_sheet("Checks")

 def nome(n, ws, ref):
     wb.defined_names[n] = DefinedName(n, attr_text=f"'{ws.title}'!${ref[0]}${ref[1:]}" if ref[1:].isdigit() else f"'{ws.title}'!{ref}")

 # ---------------- Inputs ----------------
 inp["A1"] = f"Ipotesi e dati - {lotto['titolo']}"; inp["A1"].font = TIT
 inp["A2"] = "Blu = dato o ipotesi inseriti a mano; nero = formula. Ogni riga dichiara la fonte."
 hdr = ["Voce", "Usato", "Unita'", "Fonte / stato"]
 r = 4
 for i, h in enumerate(hdr): inp.cell(r, 1 + i, h).font = BOLD
 righe = lotto["righe"]
 r = 5
 for n, et, v, f, u, fo in righe:
     inp.cell(r, 1, et); c = inp.cell(r, 2, v); c.font = BLU; c.number_format = f
     inp.cell(r, 3, u); inp.cell(r, 4, fo); nome(n, inp, f"B{r}"); r += 1

 # scenari
 r += 1; inp.cell(r, 1, "Scenari (driver incerti)").font = BOLD; r += 1
 inp.cell(r, 1, "Scenario attivo (1 Base, 2 Favorevole, 3 Sfavorevole)"); c = inp.cell(r, 2, 1); c.font = BLU; nome("Scenario", inp, f"B{r}"); rs = r; r += 1
 for i, h in enumerate(["Driver", "Usato", "Unita'", "Base", "Favorevole", "Sfavorevole", "Motivo"]): inp.cell(r, 1 + i, h).font = BOLD
 r += 1
 scen = lotto["scenari"]
 for n, et, u, b, up, dn, f, mot in scen:
     inp.cell(r, 1, et); inp.cell(r, 3, u)
     for j, v in enumerate([b, up, dn]):
         c = inp.cell(r, 4 + j, v); c.font = BLU; c.number_format = f
     c = inp.cell(r, 2, f"=CHOOSE(Scenario,D{r},E{r},F{r})"); c.number_format = f
     inp.cell(r, 7, mot); nome(n, inp, f"B{r}"); r += 1
 r += 1
 # derivati
 inp.cell(r, 1, "Calendario derivato").font = BOLD; r += 1
 der = [
  ("MeseDecreto", "Mese del decreto", "=MeseSaldo+MesiDecreto"),
  ("MeseLiber", "Mese della liberazione", "=MeseDecreto+MesiLiberazione"),
  ("MeseLavoriIni", "Inizio lavori", "=MeseLiber+1"),
  ("MeseLavoriFine", "Fine lavori", "=MeseLavoriIni+MesiLavori-1"),
  ("MeseVendita", "Mese della vendita", "=MeseLavoriFine+MesiVendita"),
  ("Arv", "Prezzo di vendita (uscita x superficie commerciale + accessori)", "=UscitaMq*MqComm+Accessori"),
  ("LavoriTot", "Lavori di finitura con imprevisti", "=LavoriMq*MqAlloggio*(1+Imprevisti_pct)"),
  ("TassoMese", "Tasso mensile equivalente all'IRR obiettivo", "=(1+Irr_obj)^(1/12)-1"),
 ]
 for n, et, fml in der:
     inp.cell(r, 1, et); c = inp.cell(r, 2, fml); c.number_format = EUR if n in ("Arv", "LavoriTot") else ('0.000%' if n == "TassoMese" else NUM)
     nome(n, inp, f"B{r}"); r += 1
 inp.column_dimensions["A"].width = 52; inp.column_dimensions["B"].width = 16; inp.column_dimensions["C"].width = 16; inp.column_dimensions["D"].width = 14
 for col in "EF": inp.column_dimensions[col].width = 13
 inp.column_dimensions["G"].width = 60

 # ---------------- CF (griglia mensile) ----------------
 N = 36  # mesi 0..36
 F0 = 5  # colonna del mese 0 (A etichetta, B unita', C totale, D spaziatore)
 cf["A1"] = "Flussi di cassa mensili (entrate +, uscite -), senza debito: capitale del club deal"; cf["A1"].font = TIT
 def riga(rr, et, un, fml, fmt=EUR, tot=True, bold=False):
     cf.cell(rr, 1, et).font = BOLD if bold else NERO; cf.cell(rr, 2, un)
     for w in range(N + 1):
         col = F0 + w; c = cf.cell(rr, col, fml(L(col), L(col - 1))); c.number_format = fmt
         if bold: c.font = BOLD; c.border = TOP
     if tot:
         c = cf.cell(rr, 3, f"=SUM({L(F0)}{rr}:{L(F0+N)}{rr})"); c.number_format = fmt
         if bold: c.font = BOLD
 R = {}
 rr = 3
 R["mese"] = rr; riga(rr, "Mese", "n", lambda c, p: f"=IF(ISNUMBER({p}{rr}),{p}{rr}+1,0)", NUM, False); rr += 1
 R["data"] = rr; riga(rr, "Data (fine mese)", "data", lambda c, p: f"=EOMONTH(DataAgg,{c}{R['mese']})", "dd/mm/yyyy", False); rr += 1
 m = lambda c: f"{c}${R['mese']}"
 R["f_saldo"] = rr; riga(rr, "Flag saldo", "1/0", lambda c, p: f"=1*({m(c)}=MeseSaldo)", NUM); rr += 1
 R["f_decr"] = rr; riga(rr, "Flag decreto", "1/0", lambda c, p: f"=1*({m(c)}=MeseDecreto)", NUM); rr += 1
 R["f_poss"] = rr; riga(rr, "Flag possesso (decreto - vendita)", "1/0", lambda c, p: f"=1*({m(c)}>MeseDecreto)*({m(c)}<=MeseVendita)", NUM); rr += 1
 R["f_lav"] = rr; riga(rr, "Flag lavori", "1/0", lambda c, p: f"=1*({m(c)}>=MeseLavoriIni)*({m(c)}<=MeseLavoriFine)", NUM); rr += 1
 R["f_lavini"] = rr; riga(rr, "Flag inizio lavori", "1/0", lambda c, p: f"=1*({m(c)}=MeseLavoriIni)", NUM); rr += 1
 R["f_vend"] = rr; riga(rr, "Flag vendita", "1/0", lambda c, p: f"=1*({m(c)}=MeseVendita)", NUM); rr += 1
 rr += 1
 f = lambda k, c: f"{c}${R[k]}"
 cf.cell(rr, 1, "Famiglia prezzo (lineare nel prezzo offerto)").font = BOLD; rr += 1
 R["cauz"] = rr; riga(rr, "Cauzione", "EUR", lambda c, p: f"=-Cauzione_pct*Prezzo*1*({m(c)}=0)"); rr += 1
 R["saldo"] = rr; riga(rr, "Saldo prezzo", "EUR", lambda c, p: f"=-(1-Cauzione_pct)*Prezzo*{f('f_saldo',c)}"); rr += 1
 R["reg"] = rr; riga(rr, "Imposta di registro", "EUR", lambda c, p: f"=-Registro_pct*Prezzo*{f('f_saldo',c)}"); rr += 1
 R["fam_p"] = rr; riga(rr, "Totale famiglia prezzo", "EUR", lambda c, p: f"={c}{R['cauz']}+{c}{R['saldo']}+{c}{R['reg']}", bold=True); rr += 2
 cf.cell(rr, 1, "Altri costi").font = BOLD; rr += 1
 R["fisse"] = rr; riga(rr, "Imposte fisse e spese di trasferimento", "EUR", lambda c, p: f"=-(ImposteFisse+SpeseTrasf)*{f('f_saldo',c)}"); rr += 1
 R["arr"] = rr; riga(rr, "Arretrati condominiali", "EUR", lambda c, p: f"=-ArretratiCond*{f('f_decr',c)}"); rr += 1
 R["cond"] = rr; riga(rr, "Condominio ordinario", "EUR", lambda c, p: f"=-CondAnnuo/12*{f('f_poss',c)}"); rr += 1
 R["imu"] = rr; riga(rr, "IMU", "EUR", lambda c, p: f"=-ImuAnnua/12*{f('f_poss',c)}"); rr += 1
 R["ute"] = rr; riga(rr, "Utenze e assicurazione", "EUR", lambda c, p: f"=-UtenzeMese*{f('f_poss',c)}"); rr += 1
 R["regol"] = rr; riga(rr, "Sanatoria e catasto", "EUR", lambda c, p: f"=-Regolarizz*{f('f_lavini',c)}"); rr += 1
 R["ripr"] = rr; riga(rr, lotto["etichetta_rischio"], "EUR", lambda c, p: f"=-Ripristino/MAX(1,MesiLavori)*{f('f_lav',c)}"); rr += 1
 R["lav"] = rr; riga(rr, "Lavori di finitura", "EUR", lambda c, p: f"=-LavoriTot/MAX(1,MesiLavori)*{f('f_lav',c)}"); rr += 1
 R["fam_c"] = rr; riga(rr, "Totale altri costi", "EUR", lambda c, p: f"=SUM({c}{R['fisse']}:{c}{R['lav']})", bold=True); rr += 2
 cf.cell(rr, 1, "Vendita").font = BOLD; rr += 1
 R["ven"] = rr; riga(rr, "Prezzo di vendita", "EUR", lambda c, p: f"=Arv*{f('f_vend',c)}"); rr += 1
 R["cven"] = rr; riga(rr, "Costi di vendita", "EUR", lambda c, p: f"=-Vendita_pct*Arv*{f('f_vend',c)}"); rr += 1
 R["fam_v"] = rr; riga(rr, "Totale vendita netta", "EUR", lambda c, p: f"={c}{R['ven']}+{c}{R['cven']}", bold=True); rr += 2
 R["cf"] = rr; riga(rr, "Flusso di cassa ante imposte", "EUR", lambda c, p: f"={c}{R['fam_p']}+{c}{R['fam_c']}+{c}{R['fam_v']}", bold=True); rr += 1
 R["cum"] = rr; riga(rr, "Cumulato", "EUR", lambda c, p: f"=IF(ISNUMBER({p}{rr}),{p}{rr},0)+{c}{R['cf']}", tot=False); rr += 1
 R["tax"] = rr; riga(rr, "IRES sull'utile (alla vendita)", "EUR", lambda c, p: f"=-Ires_pct*MAX(0,$C${R['cf']})*{f('f_vend',c)}"); rr += 1
 R["cfn"] = rr; riga(rr, "Flusso di cassa dopo IRES", "EUR", lambda c, p: f"={c}{R['cf']}+{c}{R['tax']}", bold=True); rr += 2
 cf.cell(rr, 1, "Club deal: flussi per socio").font = BOLD; rr += 1
 R["inv"] = rr; riga(rr, "Investitori", "EUR", lambda c, p: f"=-MIN(0,{c}{R['cfn']})*-(1-Sponsor_cap)+Outputs!$C$40*{f('f_vend',c)}"); rr += 1
 R["spo"] = rr; riga(rr, "Promotore (QRE)", "EUR", lambda c, p: f"=-MIN(0,{c}{R['cfn']})*-Sponsor_cap+Outputs!$C$41*{f('f_vend',c)}"); rr += 1
 cf.column_dimensions["A"].width = 44; cf.column_dimensions["B"].width = 8; cf.column_dimensions["C"].width = 14; cf.column_dimensions["D"].width = 2
 for w in range(N + 1): cf.column_dimensions[L(F0 + w)].width = 12
 cf.freeze_panes = cf.cell(3, F0)
 rng = lambda k: f"CF!$E${R[k]}:${L(F0+N)}${R[k]}"
 tot = lambda k: f"CF!$C${R[k]}"

 # ---------------- Outputs ----------------
 out["A1"] = "Risultati"; out["A1"].font = TIT
 o = [
  (3, "Prezzo offerto", "=Prezzo", EUR),
  (4, "Costo totale (prezzo, imposte, condominio, lavori, regolarizzazioni)", f"=-({tot('fam_p')}+{tot('fam_c')})", EUR),
  (5, "Prezzo di vendita netto", f"={tot('fam_v')}", EUR),
  (6, "Utile ante imposte", f"={tot('cf')}", EUR),
  (7, "Margine sul costo", "=C6/C4", PCT),
  (8, "Capitale impegnato (picco di cassa)", f"=-MIN({rng('cum')})", EUR),
  (9, "IRR annuo ante imposte", f"=(1+IRR({rng('cf')}))^12-1", PCT),
  (10, "Utile dopo IRES", f"={tot('cfn')}", EUR),
  (11, "IRR annuo dopo IRES", f"=(1+IRR({rng('cfn')}))^12-1", PCT),
  (12, "Multiplo del capitale dopo IRES", f"=SUMIF({rng('cfn')},\">0\")/-SUMIF({rng('cfn')},\"<0\")", '0.00"x"'),
  (13, "Durata dell'operazione", "=MeseVendita", '0" mesi"'),
  (15, "PREZZO MASSIMO D'OFFERTA", None, None),
  (16, "Costi fissi non legati al prezzo (positivi)", f"=-{tot('fam_c')}", EUR),
  (17, "Prezzo massimo per margine obiettivo", "=(C5/(1+Margine_obj)-C16)/(1+Registro_pct)", EUR),
  (18, "VAN al tasso obiettivo dei flussi non legati al prezzo", f"=SUMPRODUCT(CF!$E${R['fam_c']}:${L(F0+N)}${R['fam_c']}+CF!$E${R['fam_v']}:${L(F0+N)}${R['fam_v']},(1+TassoMese)^-CF!$E${R['mese']}:${L(F0+N)}${R['mese']})", EUR),
  (19, "VAN per 1 EUR di prezzo offerto", f"=SUMPRODUCT({rng('fam_p')},(1+TassoMese)^-CF!$E${R['mese']}:${L(F0+N)}${R['mese']})/Prezzo", '0.0000'),
  (20, "Prezzo massimo per IRR obiettivo (ante imposte)", "=-C18/C19", EUR),
  (21, "Prezzo massimo d'offerta (il piu' basso dei due)", "=MIN(C17,C20)", EUR),
  (22, "Vincolo che lega", "=IF(C17<=C20,\"margine\",\"IRR\")", None),
  (23, f"Rispetto all'offerta minima ({lotto['offerta_minima_testo']})", "=C21/OffertaMinima-1", PCT),
  (24, "Aggiudicazione attesa (base x moltiplicatore)", "=PrezzoBase*Moltiplicatore", EUR),
  (25, "Margine di sicurezza: prezzo massimo meno aggiudicazione attesa", "=C21-C24", EUR),
  (27, "PUNTO DI PAREGGIO SUL VALORE DI USCITA (al prezzo offerto)", None, None),
  (28, "Uscita per margine obiettivo", "=(C4*(1+Margine_obj)/(1-Vendita_pct)-Accessori)/MqComm", '#,##0" €/mq"'),
  (29, "Uscita per utile zero", "=(C4/(1-Vendita_pct)-Accessori)/MqComm", '#,##0" €/mq"'),
  (30, "Uscita usata nel caso attivo", "=UscitaMq", '#,##0" €/mq"'),
  (31, "Riscontro: valore della perizia / superficie commerciale", "=ValorePerizia/MqComm", '#,##0" €/mq"'),
  (33, "VERIFICA DEI PREZZI MASSIMI (ricalcolo indipendente)", None, None),
  (34, "Margine al prezzo massimo per margine", "=(C5-(C17*(1+Registro_pct)+C16))/(C17*(1+Registro_pct)+C16)", PCT),
  (35, "VAN al prezzo massimo per IRR (deve essere circa 0)", "=C18+C19*C20", EUR),
  (37, "CLUB DEAL (alla vendita, dopo IRES)", None, None),
  (38, "Capitale totale versato", f"=-SUMIF({rng('cfn')},\"<0\")", EUR),
  (39, "Preferenziale maturato investitori (semplice)", "=C38*(1-Sponsor_cap)*Pref_pct*MeseVendita/12", EUR),
  (40, "Distribuzione agli investitori", "=C44*(1-Sponsor_cap)+C45+C46*(1-Sponsor_split)", EUR),
  (41, "Distribuzione al promotore", "=C47-C40", EUR),
  (42, "IRR investitori", f"=(1+IRR({rng('inv')}))^12-1", PCT),
  (43, "IRR promotore", f"=(1+IRR({rng('spo')}))^12-1", PCT),
  (44, "Capitale restituito (pro quota; una perdita si ripartisce pro quota)", "=MIN(C47,C38)", EUR),
  (45, "Preferenziale pagato", "=MIN(C47-C44,C39)", EUR),
  (46, "Utile residuo da ripartire", "=C47-C44-C45", EUR),
  (47, "Distribuzioni totali", "=SUMIF(" + rng('cfn') + ",\">0\")", EUR),
  (49, "RISULTATI A PREZZI DIVERSI (margine ante imposte, forma chiusa)", None, None),
  (50, "Margine all'offerta minima", "=(C5-(OffertaMinima*(1+Registro_pct)+C16))/(OffertaMinima*(1+Registro_pct)+C16)", PCT),
  (51, "VAN al 12% all'offerta minima", "=C18+C19*OffertaMinima", EUR),
  (52, "Margine al prezzo base", "=(C5-(PrezzoBase*(1+Registro_pct)+C16))/(PrezzoBase*(1+Registro_pct)+C16)", PCT),
  (53, "Margine all'aggiudicazione attesa", "=(C5-(C24*(1+Registro_pct)+C16))/(C24*(1+Registro_pct)+C16)", PCT),
 ]
 for rrr, et, fml, fmt in o:
     c0 = out.cell(rrr, 1, et)
     if fml is None: c0.font = BOLD; continue
     c = out.cell(rrr, 3, fml)
     if fmt: c.number_format = fmt
 for rrr in (6, 7, 9, 11, 21, 28): out.cell(rrr, 3).font = BOLD; out.cell(rrr, 1).font = BOLD
 out.column_dimensions["A"].width = 66; out.column_dimensions["C"].width = 18

 # ---------------- Checks ----------------
 chk["A1"] = "Controlli (1 = problema)"; chk["A1"].font = TIT
 cs = [
  ("Le famiglie di costo e vendita sommano al flusso di cassa", f"=IFERROR(1*(ABS({tot('fam_p')}+{tot('fam_c')}+{tot('fam_v')}-{tot('cf')})>0.5),1)"),
  ("La vendita cade dentro la griglia di 36 mesi", "=IFERROR(1*(MeseVendita>MesiGriglia),1)"),
  ("Il calendario e' in ordine (saldo < decreto < lavori < vendita)", "=IFERROR(1*NOT(AND(MeseSaldo<MeseDecreto,MeseDecreto<MeseLavoriIni,MeseLavoriFine<MeseVendita)),1)"),
  ("Selettore dello scenario valido", "=IFERROR(1*(OR(Scenario<1,Scenario>NumScenari,Scenario<>INT(Scenario))),1)"),
  ("Il prezzo massimo per margine restituisce il margine obiettivo", "=IFERROR(1*(ABS(Outputs!C34-Margine_obj)>0.0001),1)"),
  ("Il prezzo massimo per IRR azzera il VAN", "=IFERROR(1*(ABS(Outputs!C35)>1),1)"),
  ("I flussi dei soci sommano al flusso dopo IRES", f"=IFERROR(1*(ABS(CF!C{R['inv']}+CF!C{R['spo']}-{tot('cfn')})>1),1)"),
  ("IRR ante imposte calcolato (non errore)", "=IFERROR(1*NOT(ISNUMBER(Outputs!C9)),1)"),
 ]
 rr = 3
 for et, fml in cs:
     chk.cell(rr, 1, et); chk.cell(rr, 2, fml); rr += 1
 re1, re2 = 3, rr - 1
 rr += 1; chk.cell(rr, 1, "Warnings").font = BOLD; rr += 1; rw1 = rr
 ws_ = [
  ("Prezzo offerto sopra il prezzo massimo d'offerta", "=IFERROR(1*(Prezzo>Outputs!C21),1)"),
  ("Margine sotto la soglia QRE", "=IFERROR(1*(Outputs!C7<Margine_obj),1)"),
  ("IRR dopo IRES sotto la soglia QRE", "=IFERROR(1*(Outputs!C11<Irr_obj),1)"),
  *[(t, 1) for t in lotto["avvisi"]],
 ]
 for et, fml in ws_:
     chk.cell(rr, 1, et); c = chk.cell(rr, 2, fml)
     if not isinstance(fml, str): c.font = BLU
     rr += 1
 rw2 = rr - 1
 chk.cell(rr + 1, 1, "Errori totali").font = BOLD; chk.cell(rr + 1, 2, f"=SUM(B{re1}:B{re2})").font = BOLD
 chk.cell(rr + 2, 1, "Avvisi totali").font = BOLD; chk.cell(rr + 2, 2, f"=SUM(B{rw1}:B{rw2})").font = BOLD
 nome("ErroriTot", chk, f"B{rr+1}"); nome("AvvisiTot", chk, f"B{rr+2}")
 chk.column_dimensions["A"].width = 78

 # ---------------- Cover ----------------
 cover["A1"] = f"Quantum Re - Fix & flip, lotto PVP {lotto['id']}"; cover["A1"].font = TIT
 cover["A2"] = lotto["descrizione"]
 cover["A3"] = "Segni: entrate positive, uscite negative. IRR: mensile composto ((1+IRR)^12-1). Nessun debito: capitale proprio del club deal."
 cv = [
  (5, "Scenario attivo", "=CHOOSE(Scenario,\"Base\",\"Favorevole\",\"Sfavorevole\")", None),
  (6, "Prezzo massimo d'offerta", "=Outputs!C21", EUR),
  (7, "Vincolo che lega", "=Outputs!C22", None),
  (8, f"Margine all'offerta minima ({lotto['offerta_minima_testo']})", "=Outputs!C50", PCT),
  (9, "Aggiudicazione attesa (base x moltiplicatore)", "=Outputs!C24", EUR),
  (10, "Al prezzo offerto: margine sul costo", "=Outputs!C7", PCT),
  (16, "Al prezzo offerto: IRR ante imposte", "=Outputs!C9", PCT),
  (11, "Uscita minima per margine 30% (al prezzo offerto)", "=Outputs!C28", '#,##0" €/mq"'),
  (12, "Uscita per utile zero", "=Outputs!C29", '#,##0" €/mq"'),
  (13, "Capitale impegnato", "=Outputs!C8", EUR),
  (14, "Errori nei controlli", "=ErroriTot", NUM),
  (15, "Avvisi", "=AvvisiTot", NUM),
 ]
 for rrr, et, fml, fmt in cv:
     cover.cell(rrr, 1, et); c = cover.cell(rrr, 3, fml); c.font = VERDE
     if fmt: c.number_format = fmt
 cover["A17"] = lotto["da_confermare"]
 cover["A18"] = lotto["fonti"]
 cover.column_dimensions["A"].width = 62; cover.column_dimensions["C"].width = 26

 wb.save(percorso)


if __name__ == "__main__":
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from lotti_fix_flip import LOTTI
    if len(sys.argv) != 2 or sys.argv[1] not in LOTTI:
        sys.exit(f"Uso: python3 analisi/modello_fix_flip.py <{'|'.join(LOTTI)}>")
    lotto = LOTTI[sys.argv[1]]
    destinazione = os.path.join(os.path.dirname(os.path.abspath(__file__)), lotto["file"])
    costruisci(lotto, destinazione)
    print(f"salvato {destinazione}")
